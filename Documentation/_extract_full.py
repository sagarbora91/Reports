"""Extract docx content with full Unicode (Marathi preserved)."""
import sys
import io
import json
from docx import Document

# Force UTF-8 stdout on Windows
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def extract(path, out_path):
    doc = Document(path)
    lines = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text
        style = p.style.name if p.style else ''
        if text.strip() or style.startswith('Heading'):
            lines.append(f"[P{i:04d}|{style}] {text}")
    # Tables
    for t_idx, table in enumerate(doc.tables):
        lines.append(f"\n=== TABLE {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ===")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.replace('\n', ' / ').strip() for c in row.cells]
            lines.append(f"  R{r_idx}: " + " | ".join(cells))
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"Wrote {out_path}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")

if __name__ == '__main__':
    base = r'C:\Cowork\Reports Apk\Documentation'
    extract(rf'{base}\Saagar_P1_App_Spec_v1.docx',        rf'{base}\_full_appspec.txt')
    extract(rf'{base}\Saagar_P1_Audit_Workbook_v1.docx',  rf'{base}\_full_workbook.txt')
    extract(rf'{base}\Saagar_P1_QuickRef_v1.docx',        rf'{base}\_full_quickref.txt')
